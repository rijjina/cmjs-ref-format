import sys
import json
import argparse
import os
import shutil
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

def apply_runs(paragraph, run_data_list, default_font="Calibri", default_size=Pt(10), default_bold=False):
    for r_data in run_data_list:
        run = paragraph.add_run(r_data["text"])
        run.font.name = default_font
        run.font.size = default_size
        run.bold = r_data.get("bold", False) or default_bold
        run.italic = r_data.get("italic", False)
        run.underline = r_data.get("underline", False)
        
        if r_data.get("superscript"):
            run.font.superscript = True
        if r_data.get("subscript"):
            run.font.subscript = True

def clear_document(doc):
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)
    for table in list(doc.tables):
        table._element.getparent().remove(table._element)

def format_document(parsed_data, template_path, output_path):
    doc = Document(template_path)
    clear_document(doc)
    
    # 1. Setup margins & page size
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(0.92)
        section.top_margin = Inches(0.71)
        section.bottom_margin = Inches(0.69)
    
    # Title
    if parsed_data.get("title"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.THAI_JUSTIFY
        apply_runs(p, parsed_data["title"], default_size=Pt(16), default_bold=True)
    
    # Authors
    for author_runs in parsed_data.get("authors", []):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        apply_runs(p, author_runs, default_size=Pt(10), default_bold=True)
        
    # Affiliations
    for affil_runs in parsed_data.get("affiliations", []):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        apply_runs(p, affil_runs, default_size=Pt(10))
        
    # Corresponding author
    if parsed_data.get("corresponding_author"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        apply_runs(p, parsed_data["corresponding_author"], default_size=Pt(10))
        
    # Sections
    for section in parsed_data.get("sections", []):
        sec_name = section["section_name"]
        
        # Add heading
        p = doc.add_paragraph()
        if sec_name.upper() in ["ABSTRACT", "HIGHLIGHTS", "INTRODUCTION", "1. INTRODUCTION", "MATERIALS AND METHODS", "2. MATERIALS AND METHODS", "RESULTS AND DISCUSSION", "3. RESULTS AND DISCUSSION", "CONCLUSIONS", "4. CONCLUSIONS", "ACKNOWLEDGEMENTS", "AUTHOR CONTRIBUTIONS", "CONFLICT OF INTEREST STATEMENT", "DECLARATION OF USE OF GENERATIVE AI", "ETHICAL GUIDELINES", "FUNDING", "REFERENCES"] or sec_name.upper().startswith("KEYWORDS"):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(sec_name)
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            run.bold = True
        else:
            # Subheading
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(sec_name)
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            run.bold = True

        # Add paragraphs
        for p_runs in section.get("paragraphs", []):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            apply_runs(p, p_runs, default_size=Pt(10))
            
    doc.save(output_path)

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Format article according to CMJS template.")
    parser.add_argument("json_input", help="Path to parsed JSON")
    parser.add_argument("output_dir", help="Directory to save the formatted document")
    parser.add_argument("--template", help="Path to template .docx", default=None)
    
    args = parser.parse_args()
    
    if not os.path.exists(args.json_input):
        print(f"Error: File {args.json_input} does not exist.", file=sys.stderr)
        sys.exit(1)
        
    if not os.path.exists(args.output_dir):
        os.makedirs(args.output_dir, exist_ok=True)
        
    template_path = args.template
    if not template_path:
        # Default to resources relative to this script
        script_dir = os.path.dirname(os.path.abspath(__file__))
        template_path = os.path.join(script_dir, "..", "resources", "cmjs-template.docx")
        
    if not os.path.exists(template_path):
        print(f"Error: Template {template_path} not found.", file=sys.stderr)
        sys.exit(1)
        
    try:
        with open(args.json_input, 'r', encoding='utf-8') as f:
            data = json.load(f)
            
        base_name = os.path.basename(args.json_input).replace('.json', '')
        output_file = os.path.join(args.output_dir, f"{base_name}_formatted.docx")
        
        format_document(data, template_path, output_file)
        print(f"Successfully created formatted document at: {output_file}")
        
    except Exception as e:
        print(f"Error formatting document: {e}", file=sys.stderr)
        sys.exit(1)
